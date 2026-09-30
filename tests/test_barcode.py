"""Code 128 barcode batch uses the same range contract as QR bulk."""

from __future__ import annotations

import io
import zipfile

from app import app as flask_app

PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _client():
    flask_app.config.update(TESTING=True)
    return flask_app.test_client()


def test_single_barcode_png():
    client = _client()
    rv = client.post("/api/qr/single", data={"data": "1001", "label": "1001", "symbol": "barcode"})
    assert rv.status_code == 200
    assert rv.mimetype == "image/png"
    assert rv.data.startswith(PNG_MAGIC)


def test_single_barcode_rejects_non_ascii():
    client = _client()
    rv = client.post("/api/qr/single", data={"data": "bibé", "symbol": "barcode"})
    assert rv.status_code == 400


def test_barcode_batch_zip_matches_range():
    client = _client()
    rv = client.post(
        "/api/qr/batch",
        data={"start": "101", "count": "3", "symbol": "barcode", "format": "zip", "prefix": "bib_"},
    )
    assert rv.status_code == 200
    assert rv.mimetype == "application/zip"
    assert "barcode_batch_101_103.zip" in rv.headers["Content-Disposition"]
    with zipfile.ZipFile(io.BytesIO(rv.data)) as zf:
        assert zf.namelist() == [
            "bib_101.png",
            "bib_102.png",
            "bib_103.png",
            "barcode_mapping.xlsx",
        ]
        assert zf.read("bib_101.png").startswith(PNG_MAGIC)
        _assert_mapping(zf.read("barcode_mapping.xlsx"), ["101", "102", "103"])


def test_barcode_jpg_named_by_sequence():
    client = _client()
    rv = client.post(
        "/api/qr/batch",
        data={
            "start": "789765",
            "count": "2",
            "symbol": "barcode",
            "format": "zip_jpg",
            "name_by": "sequence",
        },
    )
    assert rv.status_code == 200
    assert "barcode_batch_789765_789766_jpg.zip" in rv.headers["Content-Disposition"]
    with zipfile.ZipFile(io.BytesIO(rv.data)) as zf:
        assert zf.namelist()[:2] == ["1.jpg", "2.jpg"]
        jpg = zf.read("1.jpg")
        assert jpg[:3] == b"\xff\xd8\xff"
        assert len(jpg) > 8000
        _assert_mapping(zf.read("barcode_mapping.xlsx"), ["789765", "789766"])


def _assert_mapping(raw: bytes, barcode_ids: list[str]) -> None:
    import openpyxl

    sheet = openpyxl.load_workbook(io.BytesIO(raw)).active
    assert sheet["A1"].value == "Sequence"
    assert sheet["B1"].value == "Barcode ID"
    for index, barcode_id in enumerate(barcode_ids, start=2):
        assert sheet.cell(index, 1).value == index - 1
        assert str(sheet.cell(index, 2).value) == barcode_id


def test_barcode_batch_pdf_single_is_pdf():
    client = _client()
    rv = client.post(
        "/api/qr/batch",
        data={"start": "1", "count": "2", "symbol": "barcode", "format": "pdf_single"},
    )
    assert rv.status_code == 200
    assert rv.data.startswith(b"%PDF")


def test_barcode_batch_svg_zip():
    client = _client()
    rv = client.post(
        "/api/qr/batch",
        data={"start": "5", "count": "1", "symbol": "barcode", "format": "zip_svg"},
    )
    assert rv.status_code == 200
    with zipfile.ZipFile(io.BytesIO(rv.data)) as zf:
        svg = zf.read("5.svg")
    assert b"<svg" in svg


def test_index_offers_barcode_on_batch():
    client = _client()
    body = client.get("/").data
    assert b'name="symbol"' in body
    assert b"Barcode (Code 128)" in body
